from __future__ import annotations

import os
import re
import csv
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from importlib.util import find_spec
from pathlib import Path

from PIL import Image
import fitz
import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from PySide6.QtCore import QThread, Signal
from docx import Document as WordDocument

from utils.image_utils import cv_to_pil, open_image_file, pil_to_cv


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".bmp", ".tif", ".tiff"}
OFFICE_SUFFIXES = {".doc", ".docx", ".odt", ".xls", ".xlsx", ".ods", ".ppt", ".pptx"}
PRESENTATION_SUFFIXES = {".ppt", ".pptx"}
MARKUP_SUFFIXES = {".md", ".html", ".htm", ".txt"}
DEFAULT_PDF_PASSWORD = "jianji123"


class DocumentTaskError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message

    def __str__(self) -> str:
        return f"[{self.code}] {self.message}"


@dataclass(frozen=True)
class DocumentTask:
    task_id: int
    module_key: str
    module_label: str
    action_text: str
    source_path: str
    task_kind: str
    related_paths: tuple[str, ...] = ()
    password: str = ""
    text_payload: str = ""


def detect_document_backends() -> dict[str, str]:
    return {
        "libreoffice": shutil.which("libreoffice") or shutil.which("soffice") or "",
        "pandoc": shutil.which("pandoc") or "",
        "pdftoppm": shutil.which("pdftoppm") or "",
        "pdfinfo": shutil.which("pdfinfo") or "",
        "tesseract": shutil.which("tesseract") or "",
        "ffmpeg": shutil.which("ffmpeg") or "",
        "pymupdf": "installed" if find_spec("fitz") else "",
        "pypdf": "installed" if find_spec("pypdf") else "",
        "python_docx": "installed" if find_spec("docx") else "",
        "openpyxl": "installed" if find_spec("openpyxl") else "",
    }


def build_document_output_root(base_dir: Path | None = None) -> Path:
    output_root = base_dir or (Path.cwd() / "output" / "document_workbench")
    output_root.mkdir(parents=True, exist_ok=True)
    return output_root


def format_task_error(exc: Exception) -> tuple[str, str]:
    if isinstance(exc, DocumentTaskError):
        return exc.code, exc.message
    return "TASK-UNEXPECTED", str(exc)


def parse_region_payload_blocks(payload: str, *, strict: bool = True) -> list[tuple[int, int, int, int]]:
    text = payload.strip()
    if not text:
        if strict:
            raise RuntimeError("请提供区域参数，格式为 x,y,w,h，例如 40,30,120,60。")
        return []

    rects: list[tuple[int, int, int, int]] = []
    for block in text.split(";"):
        normalized = block.strip().replace("，", ",").replace(" ", "")
        if not normalized:
            continue
        parts = normalized.split(",")
        if len(parts) != 4:
            if strict:
                raise RuntimeError("区域参数格式无效，应为 x,y,w,h 或多组以 ; 分隔。")
            continue
        try:
            x, y, w, h = [int(part) for part in parts]
        except ValueError as exc:
            if strict:
                raise RuntimeError("区域参数必须是整数，例如 40,30,120,60。") from exc
            continue
        rects.append((x, y, w, h))

    if strict and not rects:
        raise RuntimeError("没有解析到有效区域。")
    return rects


def resolve_task_kind(module_key: str, source_path: str, action_text: str) -> tuple[str | None, str]:
    suffix = Path(source_path).suffix.lower()

    if module_key == "image":
        if suffix in IMAGE_SUFFIXES and action_text == "pdf":
            return "image_to_pdf", ""
        if suffix in IMAGE_SUFFIXES and action_text == "png":
            return "image_to_png", ""
        if suffix in IMAGE_SUFFIXES and action_text == "jpg":
            return "image_to_jpg", ""
        if suffix in IMAGE_SUFFIXES and action_text == "压缩":
            return "image_compress", ""
        if suffix in IMAGE_SUFFIXES and action_text == "裁剪":
            return "image_crop", ""
        if suffix in IMAGE_SUFFIXES and action_text == "缩放":
            return "image_resize", ""
        if suffix in IMAGE_SUFFIXES and action_text == "批量处理":
            return "image_batch", ""
        if suffix == ".pdf" and action_text == "png":
            return "pdf_to_png", ""
        if suffix == ".pdf" and action_text == "jpg":
            return "pdf_to_jpg", ""
        return None, "当前图片工具已接入图片转 PDF、PNG/JPG、压缩、裁剪、缩放、批量处理，以及 PDF 转 PNG/JPG。"

    if module_key == "pdf":
        if suffix == ".pdf" and action_text == "转图片":
            return "pdf_to_png", ""
        if suffix == ".pdf" and action_text == "拆分":
            return "pdf_split", ""
        if suffix == ".pdf" and action_text == "OCR识别":
            return "pdf_ocr_txt", ""
        if suffix == ".pdf" and action_text == "合并":
            return "pdf_merge", ""
        if suffix == ".pdf" and action_text == "压缩":
            return "pdf_compress", ""
        if suffix == ".pdf" and action_text == "旋转":
            return "pdf_rotate", ""
        if suffix == ".pdf" and action_text == "加密":
            return "pdf_encrypt", ""
        if suffix == ".pdf" and action_text == "解密":
            return "pdf_decrypt", ""
        if suffix == ".pdf" and action_text == "水印处理":
            return "pdf_add_text_watermark", ""
        return None, "当前 PDF 工具已接入转图片、拆分、合并、OCR、压缩、旋转、加解密。"

    if module_key == "ocr":
        if suffix == ".pdf" and action_text == "txt":
            return "pdf_ocr_txt", ""
        if suffix == ".pdf" and action_text == "docx":
            return "pdf_ocr_docx", ""
        if suffix == ".pdf" and action_text == "可搜索 PDF":
            return "pdf_ocr_searchable_pdf", ""
        if suffix == ".pdf" and action_text == "xlsx":
            return "pdf_ocr_xlsx", ""
        if suffix in IMAGE_SUFFIXES and action_text == "txt":
            return "image_ocr_txt", ""
        if suffix in IMAGE_SUFFIXES and action_text == "docx":
            return "image_ocr_docx", ""
        if suffix in IMAGE_SUFFIXES and action_text == "可搜索 PDF":
            return "image_ocr_searchable_pdf", ""
        if suffix in IMAGE_SUFFIXES and action_text == "xlsx":
            return "image_ocr_xlsx", ""
        return None, "当前 OCR 已接入图片/PDF 转 TXT、DOCX、可搜索 PDF、XLSX。"

    if module_key == "document":
        if suffix in OFFICE_SUFFIXES and action_text == "pdf":
            return "office_to_pdf", ""
        if suffix in MARKUP_SUFFIXES and action_text == "pdf":
            return "markup_to_pdf", ""
        if suffix in PRESENTATION_SUFFIXES and action_text == "图片":
            return "presentation_to_images", ""
        if suffix in MARKUP_SUFFIXES and action_text == "docx":
            return "markup_to_docx", ""
        if suffix in MARKUP_SUFFIXES and action_text == "epub":
            return "markup_to_epub", ""
        if suffix in MARKUP_SUFFIXES and action_text == "html":
            return "markup_to_html", ""
        if suffix in MARKUP_SUFFIXES and action_text == "md":
            return "text_to_markdown", ""
        if suffix in MARKUP_SUFFIXES and action_text == "txt":
            return "document_to_txt", ""
        if suffix in {".xlsx", ".xls", ".ods"} and action_text == "csv":
            return "spreadsheet_to_csv", ""
        if suffix in {".ppt", ".pptx"} and action_text == "图片":
            return "presentation_to_images", ""
        return None, "当前文档转换已接入 Office 转 PDF、PPT 转图片、Markdown/HTML/TXT 转 DOCX/EPUB/HTML/MD/TXT、表格转 CSV。"

    if module_key == "watermark":
        if suffix == ".pdf" and action_text == "文本水印删除":
            return "pdf_remove_text_watermark", ""
        if suffix in IMAGE_SUFFIXES and action_text == "文本水印删除":
            return "image_remove_text_watermark", ""
        if suffix in IMAGE_SUFFIXES and action_text == "图片修复":
            return "image_inpaint_remove", ""
        if suffix in IMAGE_SUFFIXES and action_text == "背景填充":
            return "image_background_fill", ""
        return None, "当前去水印模块已接入 PDF 文本水印删除、图片修复和背景填充。"

    return None, "当前模块的真实执行器尚未接入。"


class DocumentTaskRunner(QThread):
    task_state_changed = Signal(int, str, str, str)
    batch_finished = Signal(int, int, int)

    def __init__(self, tasks: list[DocumentTask], output_root: Path) -> None:
        super().__init__()
        self.tasks = tasks
        self.output_root = build_document_output_root(output_root)
        self.backends = detect_document_backends()

    def run(self) -> None:
        success_count = 0
        failed_count = 0
        total = len(self.tasks)

        for task in self.tasks:
            source_name = Path(task.source_path).name
            self.task_state_changed.emit(task.task_id, "running", "执行中", f"开始处理：{source_name}")
            try:
                outputs = self._execute_task(task)
            except Exception as exc:
                failed_count += 1
                code, message = format_task_error(exc)
                self.task_state_changed.emit(task.task_id, "failed", "失败", f"[{code}] {message}")
                continue

            success_count += 1
            self.task_state_changed.emit(task.task_id, "success", "成功", self._format_output_details(outputs))

        self.batch_finished.emit(total, success_count, failed_count)

    def _execute_task(self, task: DocumentTask) -> list[Path]:
        source_path = Path(task.source_path)
        if not source_path.exists():
            raise DocumentTaskError("TASK-NOT-FOUND", f"源文件不存在：{source_path}")

        if task.task_kind == "image_to_pdf":
            return [self._convert_image_to_pdf(source_path)]
        if task.task_kind == "image_to_png":
            return [self._convert_image_format(source_path, ".png")]
        if task.task_kind == "image_to_jpg":
            return [self._convert_image_format(source_path, ".jpg")]
        if task.task_kind == "pdf_to_png":
            return self._convert_pdf_to_images(source_path, image_format="png")
        if task.task_kind == "pdf_to_jpg":
            return self._convert_pdf_to_images(source_path, image_format="jpg")
        if task.task_kind == "office_to_pdf":
            return [self._convert_office_to_pdf(source_path)]
        if task.task_kind == "markup_to_docx":
            return [self._convert_markup_to_docx(source_path)]
        if task.task_kind == "markup_to_pdf":
            return [self._convert_markup_to_pdf(source_path)]
        if task.task_kind == "presentation_to_images":
            return self._convert_presentation_to_images(source_path)
        if task.task_kind == "pdf_split":
            return self._split_pdf(source_path)
        if task.task_kind == "pdf_merge":
            return [self._merge_related_pdfs(task)]
        if task.task_kind == "pdf_ocr_txt":
            return [self._ocr_pdf_to_text(source_path)]
        if task.task_kind == "image_ocr_txt":
            return [self._ocr_image_to_text(source_path)]
        if task.task_kind == "pdf_ocr_docx":
            return [self._ocr_pdf_to_docx(source_path)]
        if task.task_kind == "image_ocr_docx":
            return [self._ocr_image_to_docx(source_path)]
        if task.task_kind == "pdf_ocr_xlsx":
            return [self._ocr_pdf_to_xlsx(source_path)]
        if task.task_kind == "image_ocr_xlsx":
            return [self._ocr_image_to_xlsx(source_path)]
        if task.task_kind == "pdf_compress":
            return [self._compress_pdf(source_path)]
        if task.task_kind == "pdf_rotate":
            return [self._rotate_pdf(source_path)]
        if task.task_kind == "pdf_encrypt":
            return [self._encrypt_pdf(task, source_path)]
        if task.task_kind == "pdf_decrypt":
            return [self._decrypt_pdf(task, source_path)]
        if task.task_kind == "pdf_ocr_searchable_pdf":
            return [self._ocr_pdf_to_searchable_pdf(source_path)]
        if task.task_kind == "image_ocr_searchable_pdf":
            return [self._ocr_image_to_searchable_pdf(source_path)]
        if task.task_kind == "pdf_add_text_watermark":
            return [self._add_text_watermark(task, source_path)]
        if task.task_kind == "pdf_remove_text_watermark":
            return [self._remove_text_watermark(task, source_path)]
        if task.task_kind == "image_remove_text_watermark":
            return [self._remove_text_watermark_from_image(task, source_path)]
        if task.task_kind == "image_inpaint_remove":
            return [self._image_inpaint_remove(task, source_path)]
        if task.task_kind == "image_background_fill":
            return [self._image_background_fill(task, source_path)]
        if task.task_kind == "image_compress":
            return [self._compress_image(source_path)]
        if task.task_kind == "image_crop":
            return [self._crop_image(source_path)]
        if task.task_kind == "image_resize":
            return [self._resize_image(source_path)]
        if task.task_kind == "image_batch":
            return self._batch_process_image(source_path)
        if task.task_kind == "markup_to_epub":
            return [self._convert_markup_to_epub(source_path)]
        if task.task_kind == "markup_to_html":
            return [self._convert_markup_to_html(source_path)]
        if task.task_kind == "text_to_markdown":
            return [self._convert_text_to_markdown(source_path)]
        if task.task_kind == "document_to_txt":
            return [self._convert_document_to_txt(source_path)]
        if task.task_kind == "spreadsheet_to_csv":
            return [self._convert_spreadsheet_to_csv(source_path)]
        raise DocumentTaskError("TASK-UNKNOWN-KIND", f"未识别的任务类型：{task.task_kind}")

    def _convert_image_to_pdf(self, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "image_to_pdf" / f"{source_path.stem}.pdf")
        image = open_image_file(str(source_path)).convert("RGB")
        image.save(destination, "PDF", resolution=150.0)
        return destination

    def _convert_image_format(self, source_path: Path, target_suffix: str) -> Path:
        destination = self._unique_output_path(
            self.output_root / "image_convert" / f"{source_path.stem}{target_suffix}"
        )
        image = open_image_file(str(source_path))
        if target_suffix == ".jpg":
            image.convert("RGB").save(destination, "JPEG", quality=95, optimize=True, progressive=True)
        else:
            image.save(destination, "PNG", optimize=True)
        return destination

    def _convert_pdf_to_images(self, source_path: Path, image_format: str) -> list[Path]:
        self._require_backend("pdftoppm", "未检测到 pdftoppm，无法执行 PDF 转图片。")
        output_dir = self._unique_output_dir(self.output_root / "pdf_to_images" / source_path.stem)
        prefix = output_dir / source_path.stem
        output_dir.mkdir(parents=True, exist_ok=True)

        command = [self.backends["pdftoppm"], "-r", "180"]
        if image_format == "jpg":
            command.extend(["-jpeg", "-jpegopt", "quality=95"])
        else:
            command.append("-png")
        command.extend([str(source_path), str(prefix)])
        self._run_command(command, failure_prefix="PDF 转图片失败")

        pattern = "*.jpg" if image_format == "jpg" else "*.png"
        outputs = sorted(output_dir.glob(pattern))
        if not outputs:
            raise RuntimeError("PDF 转图片已执行，但没有生成任何输出文件。")
        return outputs

    def _convert_office_to_pdf(self, source_path: Path) -> Path:
        office_binary = self._require_backend("libreoffice", "未检测到 LibreOffice，无法执行 Office 转 PDF。")
        output_dir = self._unique_output_dir(self.output_root / "office_to_pdf" / source_path.stem)
        output_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="jianji-lo-home-", dir="/tmp") as home_dir:
            cache_dir = Path(home_dir) / "cache"
            config_dir = Path(home_dir) / "config"
            cache_dir.mkdir(parents=True, exist_ok=True)
            config_dir.mkdir(parents=True, exist_ok=True)
            command = [
                office_binary,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nolockcheck",
                "--nofirststartwizard",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(source_path),
            ]
            self._run_command(
                command,
                failure_prefix="Office 转 PDF 失败",
                extra_env={
                    "HOME": home_dir,
                    "XDG_CACHE_HOME": str(cache_dir),
                    "XDG_CONFIG_HOME": str(config_dir),
                },
                timeout=120,
            )
        outputs = sorted(output_dir.glob("*.pdf"))
        if not outputs:
            raise RuntimeError("LibreOffice 已执行，但没有生成 PDF 文件。")
        return outputs[0]

    def _convert_markup_to_docx(self, source_path: Path) -> Path:
        pandoc_binary = self._require_backend("pandoc", "未检测到 Pandoc，无法执行文档转换。")
        destination = self._unique_output_path(self.output_root / "markup_to_docx" / f"{source_path.stem}.docx")
        command = [pandoc_binary, str(source_path), "-o", str(destination)]
        self._run_command(command, failure_prefix="文档转 DOCX 失败")
        if not destination.exists():
            raise RuntimeError("Pandoc 已执行，但没有生成 DOCX 文件。")
        return destination

    def _convert_markup_to_pdf(self, source_path: Path) -> Path:
        temp_docx = self._convert_markup_to_docx(source_path)
        return self._convert_office_to_pdf(temp_docx)

    def _convert_presentation_to_images(self, source_path: Path) -> list[Path]:
        pdf_path = self._convert_office_to_pdf(source_path)
        return self._convert_pdf_to_images(pdf_path, image_format="png")

    def _convert_markup_to_epub(self, source_path: Path) -> Path:
        pandoc_binary = self._require_backend("pandoc", "未检测到 Pandoc，无法执行 EPUB 转换。")
        destination = self._unique_output_path(self.output_root / "markup_to_epub" / f"{source_path.stem}.epub")
        self._run_command([pandoc_binary, str(source_path), "-o", str(destination)], failure_prefix="文档转 EPUB 失败")
        return destination

    def _convert_markup_to_html(self, source_path: Path) -> Path:
        pandoc_binary = self._require_backend("pandoc", "未检测到 Pandoc，无法执行 HTML 转换。")
        destination = self._unique_output_path(self.output_root / "markup_to_html" / f"{source_path.stem}.html")
        self._run_command([pandoc_binary, str(source_path), "-o", str(destination)], failure_prefix="文档转 HTML 失败")
        return destination

    def _convert_text_to_markdown(self, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "text_to_markdown" / f"{source_path.stem}.md")
        content = source_path.read_text(encoding="utf-8", errors="ignore")
        destination.write_text(content, encoding="utf-8")
        return destination

    def _convert_document_to_txt(self, source_path: Path) -> Path:
        suffix = source_path.suffix.lower()
        destination = self._unique_output_path(self.output_root / "document_to_txt" / f"{source_path.stem}.txt")
        if suffix in {".md", ".txt"}:
            destination.write_text(source_path.read_text(encoding="utf-8", errors="ignore"), encoding="utf-8")
            return destination
        if suffix in {".html", ".htm"}:
            html = source_path.read_text(encoding="utf-8", errors="ignore")
            text = re.sub(r"<[^>]+>", " ", html)
            destination.write_text(re.sub(r"\s+", " ", text).strip() + "\n", encoding="utf-8")
            return destination
        raise DocumentTaskError("TASK-TXT-UNSUPPORTED", f"当前不支持将 `{suffix}` 直接转为 TXT。")

    def _convert_spreadsheet_to_csv(self, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "spreadsheet_to_csv" / f"{source_path.stem}.csv")
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(list(row))
        workbook.close()
        return destination

    def _split_pdf(self, source_path: Path) -> list[Path]:
        reader = PdfReader(str(source_path))
        if len(reader.pages) <= 1:
            destination = self._unique_output_path(self.output_root / "pdf_split" / f"{source_path.stem}_page_001.pdf")
            writer = PdfWriter()
            writer.add_page(reader.pages[0])
            with destination.open("wb") as handle:
                writer.write(handle)
            return [destination]
        output_dir = self._unique_output_dir(self.output_root / "pdf_split" / source_path.stem)
        output_dir.mkdir(parents=True, exist_ok=True)
        outputs: list[Path] = []
        for index, page in enumerate(reader.pages, start=1):
            writer = PdfWriter()
            writer.add_page(page)
            destination = output_dir / f"{source_path.stem}_page_{index:03d}.pdf"
            with destination.open("wb") as handle:
                writer.write(handle)
            outputs.append(destination)
        return outputs

    def _merge_related_pdfs(self, task: DocumentTask) -> Path:
        if task.related_paths:
            sibling_paths = [Path(path) for path in task.related_paths]
        else:
            source_path = Path(task.source_path)
            sibling_paths = sorted(path for path in source_path.parent.glob("*.pdf") if path.is_file())
        if len(sibling_paths) < 2:
            raise RuntimeError("至少需要 2 个 PDF 文件才适合执行合并。")
        writer = PdfWriter()
        for pdf_path in sibling_paths:
            reader = PdfReader(str(pdf_path))
            for page in reader.pages:
                writer.add_page(page)
        first_source = sibling_paths[0]
        output_dir = self._unique_output_dir(self.output_root / "pdf_merge" / (first_source.parent.name or first_source.stem))
        output_dir.mkdir(parents=True, exist_ok=True)
        destination = output_dir / f"{first_source.parent.name or first_source.stem}_merged.pdf"
        with destination.open("wb") as handle:
            writer.write(handle)
        return destination

    def _ocr_pdf_to_text(self, source_path: Path) -> Path:
        text_path = self._unique_output_path(self.output_root / "ocr_text" / f"{source_path.stem}.txt")

        direct_text = []
        document = fitz.open(source_path)
        for page in document:
            direct_text.append(page.get_text("text").strip())
        merged_direct_text = "\n\n".join(part for part in direct_text if part)
        if merged_direct_text.strip():
            text_path.write_text(merged_direct_text.strip() + "\n", encoding="utf-8")
            return text_path

        tesseract_binary = self._require_backend("tesseract", "未检测到 Tesseract，无法执行 OCR。")
        image_paths = self._convert_pdf_to_images(source_path, image_format="png")
        chunks = [self._run_tesseract(image_path, tesseract_binary) for image_path in image_paths]
        text_path.write_text("\n\n".join(chunk for chunk in chunks if chunk).strip() + "\n", encoding="utf-8")
        return text_path

    def _ocr_image_to_text(self, source_path: Path) -> Path:
        tesseract_binary = self._require_backend("tesseract", "未检测到 Tesseract，无法执行 OCR。")
        text_path = self._unique_output_path(self.output_root / "ocr_text" / f"{source_path.stem}.txt")
        content = self._run_tesseract(source_path, tesseract_binary)
        text_path.write_text(content.strip() + "\n", encoding="utf-8")
        return text_path

    def _ocr_pdf_to_docx(self, source_path: Path) -> Path:
        text_path = self._ocr_pdf_to_text(source_path)
        return self._text_file_to_docx(text_path, source_path.stem)

    def _ocr_image_to_docx(self, source_path: Path) -> Path:
        text_path = self._ocr_image_to_text(source_path)
        return self._text_file_to_docx(text_path, source_path.stem)

    def _text_file_to_docx(self, text_path: Path, stem: str) -> Path:
        content = text_path.read_text(encoding="utf-8").strip()
        destination = self._unique_output_path(self.output_root / "ocr_docx" / f"{stem}.docx")
        document = WordDocument()
        document.add_heading(f"OCR 结果 - {stem}", level=1)
        if content:
            for block in content.split("\n\n"):
                document.add_paragraph(block.strip())
        else:
            document.add_paragraph("OCR 未识别到可输出文本。")
        document.save(destination)
        return destination

    def _ocr_pdf_to_xlsx(self, source_path: Path) -> Path:
        text_path = self._ocr_pdf_to_text(source_path)
        return self._text_file_to_xlsx(text_path, source_path.stem)

    def _ocr_image_to_xlsx(self, source_path: Path) -> Path:
        text_path = self._ocr_image_to_text(source_path)
        return self._text_file_to_xlsx(text_path, source_path.stem)

    def _text_file_to_xlsx(self, text_path: Path, stem: str) -> Path:
        destination = self._unique_output_path(self.output_root / "ocr_xlsx" / f"{stem}.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OCR"
        content = text_path.read_text(encoding="utf-8", errors="ignore").splitlines()
        for index, line in enumerate(content, start=1):
            sheet.cell(row=index, column=1, value=line)
        if not content:
            sheet.cell(row=1, column=1, value="OCR 未识别到可输出文本。")
        workbook.save(destination)
        return destination

    def _compress_pdf(self, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "pdf_compress" / f"{source_path.stem}_compressed.pdf")
        document = fitz.open(source_path)
        document.save(destination, garbage=4, deflate=True, clean=True)
        document.close()
        return destination

    def _rotate_pdf(self, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "pdf_rotate" / f"{source_path.stem}_rotated.pdf")
        reader = PdfReader(str(source_path))
        writer = PdfWriter()
        for page in reader.pages:
            page.rotate(90)
            writer.add_page(page)
        with destination.open("wb") as handle:
            writer.write(handle)
        return destination

    def _encrypt_pdf(self, task: DocumentTask, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "pdf_encrypt" / f"{source_path.stem}_encrypted.pdf")
        reader = PdfReader(str(source_path))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        password = self._resolve_task_password(task)
        writer.encrypt(password)
        with destination.open("wb") as handle:
            writer.write(handle)
        return destination

    def _decrypt_pdf(self, task: DocumentTask, source_path: Path) -> Path:
        destination = self._unique_output_path(self.output_root / "pdf_decrypt" / f"{source_path.stem}_decrypted.pdf")
        reader = PdfReader(str(source_path))
        if reader.is_encrypted:
            password = self._resolve_task_password(task)
            status = reader.decrypt(password)
            if status == 0:
                raise RuntimeError("PDF 解密失败：提供的密码无效。")
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        with destination.open("wb") as handle:
            writer.write(handle)
        return destination

    def _ocr_pdf_to_searchable_pdf(self, source_path: Path) -> Path:
        document = fitz.open(source_path)
        has_text_layer = any(page.get_text("text").strip() for page in document)
        document.close()
        if has_text_layer:
            destination = self._unique_output_path(
                self.output_root / "ocr_searchable_pdf" / f"{source_path.stem}_searchable.pdf"
            )
            shutil.copyfile(source_path, destination)
            return destination

        image_paths = self._convert_pdf_to_images(source_path, image_format="png")
        pdf_chunks = [self._tesseract_image_to_pdf(image_path) for image_path in image_paths]
        return self._merge_generated_pdfs(pdf_chunks, source_path.stem)

    def _ocr_image_to_searchable_pdf(self, source_path: Path) -> Path:
        return self._tesseract_image_to_pdf(source_path, output_stem=source_path.stem)

    def _tesseract_image_to_pdf(self, image_path: Path, output_stem: str | None = None) -> Path:
        tesseract_binary = self._require_backend("tesseract", "未检测到 Tesseract，无法执行 OCR。")
        stem = output_stem or image_path.stem
        destination = self._unique_output_path(self.output_root / "ocr_searchable_pdf" / f"{stem}.pdf")
        output_base = destination.with_suffix("")
        command = [tesseract_binary, str(image_path), str(output_base), "-l", "chi_sim+eng", "pdf"]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode == 0 and destination.exists():
            return destination
        stderr = (result.stderr or result.stdout).strip()
        if stderr:
            raise RuntimeError(f"生成可搜索 PDF 失败：{stderr}")
        raise RuntimeError(f"生成可搜索 PDF 失败，退出码 {result.returncode}")

    def _merge_generated_pdfs(self, pdf_paths: list[Path], stem: str) -> Path:
        destination = self._unique_output_path(self.output_root / "ocr_searchable_pdf" / f"{stem}_searchable.pdf")
        writer = PdfWriter()
        for pdf_path in pdf_paths:
            reader = PdfReader(str(pdf_path))
            for page in reader.pages:
                writer.add_page(page)
        with destination.open("wb") as handle:
            writer.write(handle)
        return destination

    def _add_text_watermark(self, task: DocumentTask, source_path: Path) -> Path:
        watermark_text = task.text_payload.strip() or "CONFIDENTIAL"
        destination = self._unique_output_path(self.output_root / "pdf_watermark" / f"{source_path.stem}_watermarked.pdf")
        document = fitz.open(source_path)
        for page in document:
            rect = page.rect
            box = fitz.Rect(rect.width * 0.15, rect.height * 0.44, rect.width * 0.85, rect.height * 0.56)
            annot = page.add_freetext_annot(
                box,
                watermark_text,
                fontsize=max(14, min(24, int(rect.width / 26))),
                text_color=(0.72, 0.72, 0.72),
                fill_color=None,
                border_color=None,
                rotate=0,
                align=fitz.TEXT_ALIGN_CENTER,
            )
            annot.set_info(content=watermark_text, title="jianji-watermark")
            annot.update(opacity=0.35)
        document.save(destination, garbage=4, deflate=True, clean=True)
        document.close()
        return destination

    def _remove_text_watermark(self, task: DocumentTask, source_path: Path) -> Path:
        keyword = task.text_payload.strip()
        if not keyword:
            raise RuntimeError("去除文本水印前请提供关键词。")
        destination = self._unique_output_path(self.output_root / "pdf_watermark" / f"{source_path.stem}_cleaned.pdf")
        document = fitz.open(source_path)
        removed = 0
        for page in document:
            rects = page.search_for(keyword)
            for rect in rects:
                page.add_redact_annot(rect, fill=(1, 1, 1))
                removed += 1
            page.apply_redactions()

            for annot in list(page.annots() or []):
                annot_text = (annot.info.get("content") or "") if annot.info else ""
                annot_title = (annot.info.get("title") or "") if annot.info else ""
                if keyword in annot_text or keyword in annot_title:
                    page.delete_annot(annot)
                    removed += 1
        if removed == 0:
            document.close()
            raise RuntimeError(f"未找到包含关键词 `{keyword}` 的文本或注释水印。")
        document.save(destination, garbage=4, deflate=True, clean=True)
        document.close()
        return destination

    def _image_inpaint_remove(self, task: DocumentTask, source_path: Path) -> Path:
        return self._process_image_inpaint(task, source_path, radius=3, algorithm=cv2.INPAINT_TELEA, suffix="cleaned")

    def _image_background_fill(self, task: DocumentTask, source_path: Path) -> Path:
        return self._process_image_inpaint(task, source_path, radius=5, algorithm=cv2.INPAINT_NS, suffix="filled")

    def _remove_text_watermark_from_image(self, task: DocumentTask, source_path: Path) -> Path:
        keyword = task.text_payload.strip()
        if not keyword:
            raise RuntimeError("去除图片文字水印前请提供关键词。")
        tesseract_binary = self._require_backend("tesseract", "未检测到 Tesseract，无法执行图片文字水印删除。")
        image = open_image_file(str(source_path)).convert("RGB")
        cv_image = pil_to_cv(image)
        mask = self._build_text_match_mask(source_path, keyword, tesseract_binary, cv_image.shape[:2])
        if int(mask.max()) == 0:
            raise RuntimeError(f"未在图片中识别到包含关键词 `{keyword}` 的文字区域。")
        result = cv2.inpaint(cv_image, mask, 3, cv2.INPAINT_TELEA)
        destination = self._unique_output_path(
            self.output_root / "image_watermark" / f"{source_path.stem}_text_cleaned{source_path.suffix.lower() or '.png'}"
        )
        cv_to_pil(result, "RGB").save(destination)
        return destination

    def _compress_image(self, source_path: Path) -> Path:
        image = open_image_file(str(source_path))
        suffix = source_path.suffix.lower()
        destination = self._unique_output_path(
            self.output_root / "image_compress" / f"{source_path.stem}_compressed{suffix or '.png'}"
        )
        if suffix in {".jpg", ".jpeg"}:
            image.convert("RGB").save(destination, quality=72, optimize=True, progressive=True)
        else:
            image.save(destination, optimize=True, compress_level=8)
        return destination

    def _crop_image(self, source_path: Path) -> Path:
        image = open_image_file(str(source_path))
        width, height = image.size
        left = width // 10
        top = height // 10
        right = max(left + 1, width - left)
        bottom = max(top + 1, height - top)
        cropped = image.crop((left, top, right, bottom))
        destination = self._unique_output_path(
            self.output_root / "image_crop" / f"{source_path.stem}_cropped{source_path.suffix.lower() or '.png'}"
        )
        cropped.save(destination)
        return destination

    def _resize_image(self, source_path: Path) -> Path:
        image = open_image_file(str(source_path))
        resized = image.resize((max(1, image.width // 2), max(1, image.height // 2)), Image.LANCZOS)
        destination = self._unique_output_path(
            self.output_root / "image_resize" / f"{source_path.stem}_resized{source_path.suffix.lower() or '.png'}"
        )
        resized.save(destination)
        return destination

    def _batch_process_image(self, source_path: Path) -> list[Path]:
        compressed = self._compress_image(source_path)
        resized = self._resize_image(source_path)
        return [compressed, resized]

    def _process_image_inpaint(
        self,
        task: DocumentTask,
        source_path: Path,
        radius: int,
        algorithm: int,
        suffix: str,
    ) -> Path:
        rects = self._parse_region_payloads(task.text_payload, source_path)
        image = open_image_file(str(source_path)).convert("RGB")
        cv_image = pil_to_cv(image)
        mask = np.zeros(cv_image.shape[:2], dtype=np.uint8)
        for x, y, w, h in rects:
            mask[y : y + h, x : x + w] = 255
        result = cv2.inpaint(cv_image, mask, radius, algorithm)
        destination = self._unique_output_path(
            self.output_root / "image_watermark" / f"{source_path.stem}_{suffix}{source_path.suffix.lower() or '.png'}"
        )
        cv_to_pil(result, "RGB").save(destination)
        return destination

    def _parse_region_payloads(self, payload: str, source_path: Path) -> list[tuple[int, int, int, int]]:
        image = open_image_file(str(source_path))
        max_w, max_h = image.size
        rects: list[tuple[int, int, int, int]] = []
        for x, y, w, h in parse_region_payload_blocks(payload):
            x = max(0, min(x, max_w - 1))
            y = max(0, min(y, max_h - 1))
            w = max(1, min(w, max_w - x))
            h = max(1, min(h, max_h - y))
            rects.append((x, y, w, h))
        return rects

    def _build_text_match_mask(
        self,
        source_path: Path,
        keyword: str,
        tesseract_binary: str,
        image_shape: tuple[int, int],
    ) -> np.ndarray:
        command = [tesseract_binary, str(source_path), "stdout", "-l", "chi_sim+eng", "tsv"]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode != 0:
            stderr = (result.stderr or result.stdout).strip()
            if stderr:
                raise RuntimeError(f"文字区域识别失败：{stderr}")
            raise RuntimeError(f"文字区域识别失败，退出码 {result.returncode}")
        mask = np.zeros(image_shape, dtype=np.uint8)
        normalized_keyword = keyword.casefold()
        for line in result.stdout.splitlines()[1:]:
            parts = line.split("\t")
            if len(parts) < 12:
                continue
            text = parts[11].strip()
            if not text or normalized_keyword not in text.casefold():
                continue
            try:
                left = int(parts[6])
                top = int(parts[7])
                width = int(parts[8])
                height = int(parts[9])
            except ValueError:
                continue
            if width <= 0 or height <= 0:
                continue
            pad = 4
            x0 = max(0, left - pad)
            y0 = max(0, top - pad)
            x1 = min(image_shape[1], left + width + pad)
            y1 = min(image_shape[0], top + height + pad)
            mask[y0:y1, x0:x1] = 255
        return mask

    def _resolve_task_password(self, task: DocumentTask) -> str:
        password = task.password.strip()
        if password:
            return password
        return DEFAULT_PDF_PASSWORD

    def _run_tesseract(self, image_path: Path, tesseract_binary: str) -> str:
        command = [tesseract_binary, str(image_path), "stdout", "-l", "chi_sim+eng"]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode == 0:
            return result.stdout.strip()
        stderr = (result.stderr or result.stdout).strip()
        if stderr:
            raise RuntimeError(f"OCR 失败：{stderr}")
        raise RuntimeError(f"OCR 失败，退出码 {result.returncode}")

    def _require_backend(self, key: str, message: str) -> str:
        binary = self.backends.get(key, "")
        if not binary:
            raise RuntimeError(message)
        return binary

    def _run_command(
        self,
        command: list[str],
        failure_prefix: str,
        extra_env: dict[str, str] | None = None,
        timeout: int | None = None,
    ) -> None:
        env = None
        if extra_env is not None:
            env = os.environ.copy()
            env.update(extra_env)
        try:
            result = subprocess.run(command, capture_output=True, text=True, env=env, timeout=timeout)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(f"{failure_prefix}：命令执行超时。") from exc
        if result.returncode == 0:
            return
        stderr = (result.stderr or result.stdout).strip()
        if stderr:
            raise RuntimeError(f"{failure_prefix}：{stderr}")
        raise RuntimeError(f"{failure_prefix}，退出码 {result.returncode}")

    def _unique_output_path(self, destination: Path) -> Path:
        if not destination.exists():
            return destination
        for index in range(2, 1000):
            candidate = destination.with_name(f"{destination.stem}_{index}{destination.suffix}")
            if not candidate.exists():
                return candidate
        raise RuntimeError(f"无法分配输出文件名：{destination.name}")

    def _unique_output_dir(self, destination: Path) -> Path:
        if not destination.exists():
            return destination
        for index in range(2, 1000):
            candidate = destination.with_name(f"{destination.name}_{index}")
            if not candidate.exists():
                return candidate
        raise RuntimeError(f"无法分配输出目录：{destination.name}")

    def _format_output_details(self, outputs: list[Path]) -> str:
        if not outputs:
            return "任务完成，但没有记录输出文件。"
        if len(outputs) == 1:
            return f"输出：{outputs[0]}"
        preview = "；".join(str(path) for path in outputs[:3])
        if len(outputs) > 3:
            preview = f"{preview}；共 {len(outputs)} 个文件"
        return f"输出：{preview}"
